"""读取大建API技术文档"""
import subprocess
import sys

# 确保安装 python-docx
subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])

from docx import Document

doc_path = r'C:\Users\poonx\Dajian_Listing_Tool\docs\运单号查询____-已融合.docx'
doc = Document(doc_path)

print("=" * 60)
print("大建API技术文档内容")
print("=" * 60)

# 输出段落
for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        print(para.text)
    if i > 150:  # 限制输出量
        break

print("\n" + "=" * 60)
print("表格内容")
print("=" * 60)

# 输出表格
for table_idx, table in enumerate(doc.tables[:5]):  # 只输出前5个表格
    print(f"\n--- 表格 {table_idx + 1} ---")
    for row in table.rows[:20]:  # 每个表格最多20行
        cells = [cell.text.strip() for cell in row.cells]
        print(" | ".join(cells))
